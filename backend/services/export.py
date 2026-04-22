from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill


def calculer_retard(date_facture: str, delai_jours: int, date_echeance: str | None = None) -> dict:
    try:
        if date_echeance:
            echeance = datetime.strptime(date_echeance, "%Y-%m-%d")
        else:
            date = datetime.strptime(date_facture, "%Y-%m-%d")
            echeance = date + timedelta(days=delai_jours)
        aujourd_hui = datetime.now()
        retard = (aujourd_hui - echeance).days
        return {
            "echeance": echeance.strftime("%Y-%m-%d"),
            "retard_jours": max(0, retard),
            "statut": "en_retard" if retard > 0 else "dans_les_delais",
        }
    except Exception:
        return {"echeance": "-", "retard_jours": 0, "statut": "inconnu"}


def calculer_penalite(montant_ttc: float, retard_jours: int, taux_bam: float = 2.25) -> float:
    taux_total = (taux_bam + 3) / 100
    penalite = montant_ttc * taux_total * (retard_jours / 365)
    return round(penalite, 2)


def _style_header(ws):
    fill = PatternFill("solid", fgColor="D9E2F3")
    font = Font(bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font


def _autosize(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 28)


def build_export_workbook(factures, ecritures, taux_bam: float = 2.25):
    wb = openpyxl.Workbook()

    journal = wb.active
    journal.title = "Journal SAGE"
    journal.append(["Numero facture", "Date", "Compte PCM", "Code tiers", "Libelle", "Debit (MAD)", "Credit (MAD)"])
    facture_map = {facture.id: facture for facture in factures}
    for ecriture in ecritures:
        facture = facture_map.get(ecriture.facture_id)
        journal.append(
            [
                facture.num_facture or facture.nom_fichier if facture else "",
                facture.date_facture if facture else "",
                ecriture.compte_pcm,
                facture.code_tiers if facture else "",
                ecriture.libelle,
                ecriture.debit or 0,
                ecriture.credit or 0,
            ]
        )
    _style_header(journal)
    _autosize(journal)

    balance = wb.create_sheet("Balance PCM")
    balance.append(["Compte PCM", "Debit total", "Credit total", "Solde"])
    grouped = {}
    for ecriture in ecritures:
        row = grouped.setdefault(ecriture.compte_pcm, {"debit": 0.0, "credit": 0.0})
        row["debit"] += float(ecriture.debit or 0)
        row["credit"] += float(ecriture.credit or 0)
    for compte, totals in sorted(grouped.items()):
        balance.append([compte, round(totals["debit"], 2), round(totals["credit"], 2), round(totals["debit"] - totals["credit"], 2)])
    _style_header(balance)
    _autosize(balance)

    tva_sheet = wb.create_sheet("TVA mensuelle")
    tva_sheet.append(["Mois", "TVA collectee", "TVA deductible", "Solde DGI"])
    monthly = {}
    for facture in factures:
        month = (facture.date_facture or "")[:7] or "inconnu"
        row = monthly.setdefault(month, {"collectee": 0.0, "deductible": 0.0})
        if (facture.type or "achat") == "vente":
            row["collectee"] += float(facture.tva or 0)
        else:
            row["deductible"] += float(facture.tva or 0)
    for month, values in sorted(monthly.items()):
        solde = round(values["collectee"] - values["deductible"], 2)
        tva_sheet.append([month, round(values["collectee"], 2), round(values["deductible"], 2), solde])
    _style_header(tva_sheet)
    _autosize(tva_sheet)

    penalites = wb.create_sheet("Penalites Loi 69-21")
    penalites.append(["Numero facture", "Fournisseur", "Date facture", "Echeance", "Retard (jours)", "Montant TTC", "Penalite"])
    for facture in factures:
        if (facture.type or "achat") == "vente":
            continue
        info = calculer_retard(facture.date_facture, 60, facture.date_echeance)
        penalty = calculer_penalite(float(facture.montant_ttc or 0), info["retard_jours"], taux_bam) if info["retard_jours"] > 0 else 0
        if penalty <= 0:
            continue
        penalites.append(
            [
                facture.num_facture or facture.nom_fichier,
                facture.fournisseur,
                facture.date_facture,
                info["echeance"],
                info["retard_jours"],
                float(facture.montant_ttc or 0),
                penalty,
            ]
        )
    _style_header(penalites)
    _autosize(penalites)

    grand_livre = wb.create_sheet("Grand Livre")
    grand_livre.append(["Compte PCM", "Date", "Numero facture", "Libelle", "Debit", "Credit"])
    for ecriture in sorted(ecritures, key=lambda item: (item.compte_pcm or "", item.id or 0)):
        facture = facture_map.get(ecriture.facture_id)
        grand_livre.append(
            [
                ecriture.compte_pcm,
                facture.date_facture if facture else "",
                facture.num_facture or facture.nom_fichier if facture else "",
                ecriture.libelle,
                ecriture.debit or 0,
                ecriture.credit or 0,
            ]
        )
    _style_header(grand_livre)
    _autosize(grand_livre)

    return wb
